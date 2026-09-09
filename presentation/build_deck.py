"""Build the LLM-MRTA methodology deck (presentation/DESIGN.md).

    python3 presentation/build_deck.py

Every figure here is traceable to docs/RESEARCH_CONTRACT.md v1.66 or a
docs/*_RESULTS.md file. No unrun numbers (CLAUDE.md).
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- design tokens (DESIGN.md) --------------------------------------------
GROUND = RGBColor(0xFB, 0xFC, 0xFB)
INK = RGBColor(0x16, 0x21, 0x1D)
MUTED = RGBColor(0x5E, 0x6D, 0x66)
LLM = RGBColor(0xA8, 0x72, 0x1F)
LLM_SOFT = RGBColor(0xF3, 0xE9, 0xD8)
DET = RGBColor(0x1C, 0x6B, 0x57)
DET_SOFT = RGBColor(0xDC, 0xEA, 0xE4)
LIMIT = RGBColor(0x8F, 0x3A, 0x1E)
RULE = RGBColor(0xD5, 0xDD, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "맑은 고딕"
MONO = "Consolas"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_W = EMU_W - 2 * MARGIN

FOOT = "계약 v1.66 · Validator 1.4 · pytest 936"

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

TOTAL = 13
_n = 0


# ---- primitives ---------------------------------------------------------
def slide():
    global _n
    _n += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = GROUND
    bg.line.fill.background()
    bg.shadow.inherit = False
    _send_back(bg)
    return s


def _send_back(shape):
    sp = shape._element
    sp.getparent().remove(sp)
    shape._element.getparent()  # noop guard
    # reinsert right after the non-shape elements (index 2 == first shape slot)
    parent = shape.part.slide.shapes._spTree
    parent.insert(2, sp)


def _no_autofit(tf):
    # keep PowerPoint from shrinking text unpredictably
    el = tf._txBody
    bodyPr = el.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        e = bodyPr.find(qn(tag))
        if e is not None:
            bodyPr.remove(e)
    bodyPr.append(el.makeelement(qn("a:noAutofit"), {}))


def text(s, x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         leading=1.12, space_after=4):
    """runs: list of paragraphs; each paragraph is a list of (str, dict)."""
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    _no_autofit(tf)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = leading
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        for txt, st in para:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = st.get("mono") and MONO or SANS
            f.size = Pt(st.get("size", 14))
            f.bold = st.get("bold", False)
            f.color.rgb = st.get("color", INK)
            if st.get("spacing"):
                _letter_spacing(r, st["spacing"])
    return box


def _letter_spacing(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(pts * 100)))


def rect(s, x, y, w, h, fill, *, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def hline(s, x, y, w, color=RULE, weight=1.0):
    t = Pt(weight)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y - t / 2), int(w), int(t))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def vline(s, x, y, h, color=RULE, weight=1.0):
    t = Pt(weight)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x - t / 2), int(y), int(t), int(h))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def chrome(s, eyebrow, title, *, title_size=30):
    hline(s, MARGIN, Inches(0.66), CONTENT_W, RULE, 1.0)
    text(s, MARGIN, Inches(0.30), CONTENT_W, Inches(0.34),
         [[(eyebrow, dict(mono=True, size=10, color=DET, bold=True, spacing=1.4))]])
    text(s, MARGIN, Inches(0.80), CONTENT_W, Inches(1.05),
         [[(title, dict(size=title_size, bold=True, color=INK))]], leading=1.06)
    # footer
    text(s, MARGIN, EMU_H - Inches(0.50), Inches(9), Inches(0.3),
         [[(FOOT, dict(mono=True, size=9, color=MUTED))]])
    text(s, EMU_W - MARGIN - Inches(1.4), EMU_H - Inches(0.50), Inches(1.4), Inches(0.3),
         [[(f"{_n:02d} / {TOTAL}", dict(mono=True, size=9, color=MUTED))]],
         align=PP_ALIGN.RIGHT)


def bullets(items, size=13.5, color=INK, gap=5):
    out = []
    for it in items:
        if isinstance(it, tuple):
            lead, rest = it
            out.append([("· ", dict(size=size, color=MUTED)),
                        (lead, dict(size=size, color=color, bold=True)),
                        (rest, dict(size=size, color=color))])
        else:
            out.append([("· ", dict(size=size, color=MUTED)),
                        (it, dict(size=size, color=color))])
    return out


def split(s, top, height, left_head, left_items, right_head, right_items, out_line):
    """Two-column amber | teal structure block."""
    midx = MARGIN + CONTENT_W / 2
    gut = Inches(0.28)
    # amber panel
    rect(s, MARGIN, top, CONTENT_W / 2 - gut / 2, height, LLM_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, MARGIN + Inches(0.26), top + Inches(0.18),
         CONTENT_W / 2 - gut, Inches(0.34),
         [[(left_head, dict(mono=True, size=10, bold=True, color=LLM, spacing=1.2))]])
    text(s, MARGIN + Inches(0.26), top + Inches(0.58),
         CONTENT_W / 2 - gut - Inches(0.2), height - Inches(0.7),
         bullets(left_items, size=12.5, color=INK), leading=1.28, space_after=7)
    # teal panel
    rx = midx + gut / 2
    rw = CONTENT_W / 2 - gut / 2
    rect(s, rx, top, rw, height, DET_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    vline(s, rx, top + Inches(0.12), height - Inches(0.24), DET, 2.0)
    text(s, rx + Inches(0.24), top + Inches(0.16), rw - Inches(0.4), Inches(0.34),
         [[(right_head, dict(mono=True, size=10, bold=True, color=DET, spacing=1.2))]])
    text(s, rx + Inches(0.24), top + Inches(0.56), rw - Inches(0.44), height - Inches(0.7),
         bullets(right_items, size=12.5, color=INK), leading=1.2, space_after=5)
    # output line
    oy = top + height + Inches(0.14)
    hline(s, MARGIN, oy, CONTENT_W, RULE, 1.0)
    text(s, MARGIN, oy + Inches(0.08), CONTENT_W, Inches(0.5),
         [[("산출 → ", dict(mono=True, size=10, bold=True, color=MUTED, spacing=1.0)),
           (out_line, dict(size=12.5, color=INK))]], leading=1.15)


def table(s, x, y, w, rows, col_w, *, tint_last=False, head_rows=1, row_h=Inches(0.42)):
    nrow, ncol = len(rows), len(rows[0])
    gt = s.shapes.add_table(nrow, ncol, x, y, w, row_h * nrow).table
    gt.first_row = False
    gt.horz_banding = False
    for j, cw in enumerate(col_w):
        gt.columns[j].width = cw
    for i, row in enumerate(rows):
        gt.rows[i].height = row_h
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.margin_left = Inches(0.10)
            c.margin_right = Inches(0.10)
            c.margin_top = Inches(0.04)
            c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = i < head_rows
            c.fill.solid()
            if is_head:
                c.fill.fore_color.rgb = INK
            elif tint_last and j == ncol - 1:
                c.fill.fore_color.rgb = DET_SOFT
            else:
                c.fill.fore_color.rgb = WHITE
            tf = c.text_frame
            tf.word_wrap = True
            _no_autofit(tf)
            p = tf.paragraphs[0]
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = str(val)
            f = r.font
            mono = (j == 0 and not is_head) or any(ch.isdigit() for ch in str(val)) and j > 0
            f.name = MONO if (is_head or (j > 0 and _looksnum(val))) else SANS
            f.size = Pt(9.5 if is_head else 10.5)
            f.bold = is_head or (tint_last and j == ncol - 1 and not is_head)
            f.color.rgb = WHITE if is_head else INK
    _strip_table_style(gt)
    return gt


def _looksnum(v):
    v = str(v).strip()
    return v and sum(ch.isdigit() for ch in v) >= len(v.replace(" ", "").replace("/", "")
                                                        .replace(".", "").replace("s", "")) / 2


def _strip_table_style(tbl):
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for child in list(tblPr):
            tblPr.remove(child)


# =========================================================================
# 01 — title
# =========================================================================
s = slide()
hline(s, MARGIN, Inches(2.05), CONTENT_W, INK, 2.0)
text(s, MARGIN, Inches(1.30), CONTENT_W, Inches(0.4),
     [[("MP4MR-INSPIRED · 재난 대응 이종 무인체계", dict(mono=True, size=11, color=DET, bold=True, spacing=1.6))]])
text(s, MARGIN, Inches(2.25), CONTENT_W, Inches(1.8),
     [[("자연어 재난 임무의", dict(size=40, bold=True, color=INK))],
      [("결정론적 검증과 이종 무인체계 할당", dict(size=40, bold=True, color=INK))]],
     leading=1.12)
text(s, MARGIN, Inches(4.35), Inches(9.4), Inches(1.4),
     [[("LLM이 생성한 임무 그래프를 또 다른 LLM Critic이 아니라 ",
        dict(size=15, color=MUTED)),
       ("결정론적 whole-graph Validator", dict(size=15, color=INK, bold=True)),
       ("로", dict(size=15, color=MUTED))],
      [("검증하고, 통과한 그래프만 CBBA와 시뮬레이터에 전달한다",
        dict(size=15, color=MUTED))]], leading=1.4)
text(s, MARGIN, EMU_H - Inches(0.95), CONTENT_W, Inches(0.6),
     [[("근거 · docs/RESEARCH_CONTRACT.md v1.66 · D-001~D-070 · "
        "docs/{P6,P8_4,P9,P12}_RESULTS.md",
        dict(mono=True, size=9, color=MUTED))]])

# =========================================================================
# 02 — motivation
# =========================================================================
s = slide()
chrome(s, "MOTIVATION", "임무 계획의 자동화, 그리고 신뢰의 문제")
text(s, MARGIN, Inches(1.95), CONTENT_W, Inches(0.9),
     [[("기존 임무 계획은 운용자가 좌표·수행시간·보상·선행조건을 ",
        dict(size=14, color=INK)),
       ("모두 수동으로", dict(size=14, color=INK, bold=True)),
       (" 입력한다. 빠르게 변하는 재난 환경에서 이는 현실적으로 어렵다",
        dict(size=14, color=INK))]], leading=1.5)

by = Inches(3.0)
for i, (q, a) in enumerate([
    ("Q1.  자연어 명령을 다중 에이전트 임무 그래프로 바꿀 수 있는가?",
     "→ LLM이 task_type · target · dependency edge 를 생성. 나머지는 결정론적 compiler."),
    ("Q2.  그 LLM 출력을 무엇으로 믿는가?",
     "→ 또 다른 LLM Critic이 아니라, 같은 후보에 항상 같은 판정을 내리는 결정론적 Validator."),
]):
    ry = by + i * Inches(1.65)
    rect(s, MARGIN, ry, CONTENT_W, Inches(1.4), WHITE, line=RULE, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    vline(s, MARGIN, ry + Inches(0.16), Inches(1.08), DET, 2.5)
    text(s, MARGIN + Inches(0.3), ry + Inches(0.20), CONTENT_W - Inches(0.6), Inches(0.5),
         [[(q, dict(size=14.5, bold=True, color=INK))]])
    text(s, MARGIN + Inches(0.3), ry + Inches(0.78), CONTENT_W - Inches(0.6), Inches(0.5),
         [[(a, dict(size=12.5, color=MUTED))]])

# =========================================================================
# 03 — related work
# =========================================================================
s = slide()
chrome(s, "RELATED WORK", "두 선행 구조와 어디가 다른가")
text(s, MARGIN, Inches(1.78), CONTENT_W, Inches(0.5),
     [[("우열을 주장하지 않는다 — 동일 조건 직접 비교를 하지 않았다. "
        "구조적 차이와 재현성 특성만", dict(size=12.5, color=MUTED))]])
rows = [
    ["", "MP4MR (ICROS 2025)", "LSMP (발표)", "본 연구"],
    ["실현가능성 검증", "LLM Critic 4단계", "운용개념 기반 검증 모듈", "결정론적 whole-graph Validator #1~#14"],
    ["동일 입력 재현성", "LLM 피드백 루프 의존", "—", "같은 후보 → 같은 판정, graph_hash 고정"],
    ["할당", "Belief Propagation", "Coverage / VRP / poly-line", "CBBA rolling READY-frontier"],
    ["실행 중 변경", "미구현 (논문이 후속과제로 명시)", "LLM이 재계획 도구 호출", "결정론적 selective release, LLM은 intent만"],
    ["모호한 명령", "—", "—", "fail-closed — 추측 대신 clarification"],
    ["시연", "ROS2 · Gazebo 3D", "군집 sim + 실 비행시험", "2D discrete-event + 네이티브 콘솔 2창"],
]
table(s, MARGIN, Inches(2.35), CONTENT_W, rows,
      [Inches(1.75), Inches(3.15), Inches(2.9), Inches(4.31)],
      tint_last=True, row_h=Inches(0.62))

# =========================================================================
# 04 — overall structure
# =========================================================================
s = slide()
chrome(s, "FRAMEWORK", "다섯 계층으로 쌓인 하나의 파이프라인")
text(s, MARGIN, Inches(1.72), CONTENT_W, Inches(0.5),
     [[("각 계층은 연구 질문 하나에 대응하고, ", dict(size=13, color=MUTED)),
       ("아래 계층이 통과해야 위 계층이 성립한다", dict(size=13, color=INK, bold=True)),
       (" — 경계는 모든 계층에서 같다", dict(size=13, color=MUTED))]])
layers = [
    ("RQ1", "자연어 명령 → 검증된 task graph", "P5 · P6"),
    ("RQ2", "검증된 graph → 이종 UAV/UGV 할당과 2D 실행", "P3 · P4 · P6.5"),
    ("RQ3", "실행 전 다중 턴 운용자 대화", "P8.0 ~ P8.4"),
    ("RQ4", "실행 중 checkpoint에서 선택적 재할당", "P9 · P10"),
    ("RQ5", "상황 발생 대응 정책의 자연어 추출", "P12"),
]
ly = Inches(2.45)
lh = Inches(0.82)
for rq, title, phase in layers:
    rect(s, MARGIN, ly, CONTENT_W, lh - Inches(0.14), WHITE, line=RULE, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, MARGIN, ly, Inches(0.92), lh - Inches(0.14), DET_SOFT,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, MARGIN, ly + Inches(0.12), Inches(0.92), Inches(0.4),
         [[(rq, dict(mono=True, size=12, bold=True, color=DET))]], align=PP_ALIGN.CENTER)
    text(s, MARGIN + Inches(1.15), ly + Inches(0.12), Inches(8.0), Inches(0.44),
         [[(title, dict(size=13.5, bold=True, color=INK))]])
    text(s, EMU_W - MARGIN - Inches(2.2), ly + Inches(0.14), Inches(2.0), Inches(0.4),
         [[(phase, dict(mono=True, size=9.5, color=MUTED))]], align=PP_ALIGN.RIGHT)
    ly += lh

# =========================================================================
# 05 — the boundary (core)
# =========================================================================
s = slide()
chrome(s, "THE BOUNDARY — 핵심", "LLM이 생성하는 것 / 결정론적 코드가 판정하는 것")
midx = MARGIN + CONTENT_W / 2
gut = Inches(0.34)
top = Inches(2.05)
h = Inches(3.75)
# amber
rect(s, MARGIN, top, CONTENT_W / 2 - gut / 2, h, LLM_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, MARGIN + Inches(0.3), top + Inches(0.24), CONTENT_W / 2 - gut, Inches(0.5),
     [[("LLM  ·  확률적 · 재시도 가능", dict(mono=True, size=11, bold=True, color=LLM, spacing=1.2))]])
text(s, MARGIN + Inches(0.3), top + Inches(0.85), CONTENT_W / 2 - gut - Inches(0.2), h - Inches(1.1),
     bullets([
         ("task 목록 ", "{task_type, target}"),
         "task 간 dependency edge",
         "5종 대화 행위 분류 + slot 추출",
         "화재 발생 시 대응 깊이 (workflow prefix)",
         "schema 실패 시 1회 repair (D-052)",
     ], size=13, gap=8), leading=1.35, space_after=9)
# teal
rx = midx + gut / 2
rw = CONTENT_W / 2 - gut / 2
rect(s, rx, top, rw, h, DET_SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
vline(s, rx, top + Inches(0.16), h - Inches(0.32), DET, 3.0)
text(s, rx + Inches(0.3), top + Inches(0.24), rw - Inches(0.5), Inches(0.5),
     [[("결정론적 코드  ·  재현 · 해시 감사", dict(mono=True, size=11, bold=True, color=DET, spacing=1.2))]])
text(s, rx + Inches(0.3), top + Inches(0.85), rw - Inches(0.55), h - Inches(1.1),
     bullets([
         "좌표 · priority · capability · duration 파생",
         "whole-graph Validator invariant #1~#14",
         "canonical MissionPatch · atomic commit / rollback",
         "CBBA 할당 · selective release 집합",
         "clarification · REJECT · UNSUPPORTED 판정",
     ], size=13, gap=8), leading=1.35, space_after=9)
text(s, MARGIN, top + h + Inches(0.16), CONTENT_W, Inches(0.4),
     [[("LLM은 Validator가 승인하지 않는 것을 아무것도 만들 수 없다 — "
        "그래서 실시간 루프에 넣어도 된다.",
        dict(size=12, color=MUTED, bold=True))]], align=PP_ALIGN.CENTER)

# =========================================================================
# 06 — RQ1 + RQ2
# =========================================================================
s = slide()
chrome(s, "계층 1 · 2 — RQ1 · RQ2", "자연어 → 검증된 graph → 이종 할당 → 2D 실행")
split(
    s, Inches(2.05), Inches(2.75),
    "LLM이 만드는 것",
    [("Step 1  ", "task 목록 {task_type, target}"),
     ("Step 2  ", "dependency edge"),
     "RQ2 계층에는 LLM 호출이 없다"],
    "결정론적 코드가 하는 것",
    ["compiler가 좌표·priority·capability를 semantic scene에서 resolve",
     "whole-graph Validator #1~#12 — 거부 시 명시적 REJECT",
     "CBBA rolling READY-frontier epoch, λ = 0.999",
     "UAV 직선 · UGV route-graph Dijkstra 이동비용",
     "SimExecutor 2D discrete-event 실행"],
    "승인된 graph + graph_hash · scene_hash · validator_version  ·  "
    "P6 실측 9/9 exact · 골든 makespan 359.84 / 257.85 s · 위반 0",
)

# =========================================================================
# 07 — RQ3
# =========================================================================
s = slide()
chrome(s, "계층 3 — RQ3", "실행 전 다중 턴 운용자 대화")
split(
    s, Inches(2.05), Inches(2.95),
    "LLM이 만드는 것",
    ["5종 대화 행위 — NEW_MISSION / REPORT_INCIDENT / UPDATE_MISSION / QUERY_STATUS / UNSUPPORTED",
     "slot 추출 — zone, 대상 표현, 대응 단계"],
    "결정론적 코드가 하는 것",
    [("fail-closed grounder  ", "— 해석 불가능한 표현은 추측 없이 clarification"),
     "canonical MissionPatch → Validator 재검증 #13 #14",
     "apply_patch atomic commit / 거부 시 원본 완전 보존",
     "후보 선택은 LLM 없이 구조화된 entity_id 로"],
    "턴별 TurnAudit — intent · grounding · patch · 전후 hash · 그 턴의 모든 모델 호출  ·  "
    "P8.4  grounder-only 12/12 · end-to-end 6/12 · 잘못된 추측 0/3",
)

# =========================================================================
# 08 — RQ4
# =========================================================================
s = slide()
chrome(s, "계층 4 — RQ4", "실행 중 checkpoint에서의 선택적 재할당")
split(
    s, Inches(2.05), Inches(2.75),
    "LLM이 만드는 것",
    ["정지 상태에서 받은 후속 명령의 intent · slot",
     "release 집합 · agent 선택에는 관여하지 않는다"],
    "결정론적 코드가 하는 것",
    ["task-completion checkpoint에서 정지 — 임의 시각 interrupt 아님",
     ("bidder-connected selective release  ",
      "— 신규 READY task와 입찰자가 겹치는 미시작 assignment만"),
     "COMPLETED · RUNNING commitment 보존, 같은 시각·위치에서 재개"],
    "CheckpointAudit · OnlineReallocationAudit — release 집합, 전후 assignment, 추가 consensus round  ·  "
    "P9  release 0 / 2 / 1 · 3정책 모두 COMPLETED · 위반 0",
)
text(s, MARGIN, Inches(6.55), CONTENT_W, Inches(0.4),
     [[("주의 — §19.3 bundle-suffix 확장은 테스트된 경로에서 한 번도 동작하지 않았다 "
        "(suffix_extra_release_count = 0, D-042).", dict(size=10.5, color=LIMIT))]])

# =========================================================================
# 09 — RQ5
# =========================================================================
s = slide()
chrome(s, "계층 5 — RQ5", "상황 발생 대응 정책의 자연어 추출")
split(
    s, Inches(2.05), Inches(2.55),
    "LLM이 만드는 것",
    ["최초 명령에서 초기 정찰 graph 와 화재 발생 시 대응 깊이 를 분리 추출",
     "실행 중 자연어 화재 신고의 zone 과 대응 단계"],
    "결정론적 코드가 하는 것",
    ["strict latent fixture가 AREA_RECON 완료 시점에 FIRE_DETECTED 1회 공개",
     "센서 · 운용자 입력이 동일한 atomic incident transaction 으로 수렴",
     "이후 경로는 RQ3 · RQ4와 완전히 동일"],
    "같은 scene, 문장 하나 바꿈 → policy · workflow prefix · 경로가 달라진다  "
    "(held-out paraphrase 8/8 exact)",
)
# counterfactual mini-table
cf = [
    ["명령의 대응 단계", "대응 task", "결과"],
    ["열화상 재확인까지만", "1", "정찰 후 THERMAL_RECON"],
    ["지상 점검 단계까지", "3", "+2 UGV task, 경로 재계산"],
    ["지상 진압 단계까지", "4", "+2 selective release, 경로 2개 변경"],
]
table(s, MARGIN, Inches(5.30), Inches(9.2), cf,
      [Inches(3.2), Inches(1.5), Inches(4.5)], row_h=Inches(0.36))

# =========================================================================
# 10 — results
# =========================================================================
s = slide()
chrome(s, "EVIDENCE", "계층별 실측 결과")
text(s, MARGIN, Inches(1.72), CONTENT_W, Inches(0.5),
     [[("전부 ", dict(size=12, color=MUTED)),
       ("gpt-5-mini-2025-08-07", dict(mono=True, size=11, color=INK)),
       (" 단일 스냅샷 · 소표본 — 모든 gold는 해당 LLM 호출 전에 커밋했고 집계는 원자료 JSON에서 독립 재계산",
        dict(size=12, color=MUTED))]])
cards = [
    ("RQ1 · P6", "task graph 생성", ["승인  9 / 9", "task P/R  1.00 / 1.00", "exact match  9 / 9", "repair  0회"]),
    ("RQ2 · P3 · P4", "이종 할당과 실행", ["plan makespan  359.84 s", "exec makespan  257.85 s", "capability 위반  0", "precedence 위반  0"]),
    ("P6.5", "통합 실행", ["demo_pass  3 / 3", "A1 graph_hash  일치", "골든 재현  일치"]),
    ("RQ3 · P8.4", "운용자 대화  12d / 36t", ["grounder-only  12 / 12", "end-to-end live  6 / 12", "clarification recall  3 / 3", "잘못된 추측  0 / 3"]),
    ("RQ4 · P9", "선택적 재할당  3정책", ["release 수  0 / 2 / 1", "makespan  453.88 s 동일", "종료  3× COMPLETED", "위반  0 / 0"]),
    ("RQ5 · P12", "대응 정책 counterfactual", ["첫 Live  4 / 8", "held-out  8 / 8", "실행 case  전부 COMPLETED", "위반  0 / 0"]),
]
cw = (CONTENT_W - Inches(0.5)) / 3
ch = Inches(2.15)
for i, (tag, title, metrics) in enumerate(cards):
    cx = MARGIN + (i % 3) * (cw + Inches(0.25))
    cy = Inches(2.35) + (i // 3) * (ch + Inches(0.25))
    rect(s, cx, cy, cw, ch, WHITE, line=RULE, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, cx + Inches(0.2), cy + Inches(0.14), cw - Inches(0.4), Inches(0.3),
         [[(tag, dict(mono=True, size=9, bold=True, color=DET, spacing=0.8))]])
    text(s, cx + Inches(0.2), cy + Inches(0.44), cw - Inches(0.4), Inches(0.34),
         [[(title, dict(size=12, bold=True, color=INK))]])
    text(s, cx + Inches(0.2), cy + Inches(0.82), cw - Inches(0.4), ch - Inches(0.95),
         [[(m.split("  ")[0] + "  ", dict(size=10.5, color=MUTED)),
           (m.split("  ", 1)[1] if "  " in m else "", dict(mono=True, size=10.5, color=INK))]
          for m in metrics], leading=1.3, space_after=3)

# =========================================================================
# 11 — reproducibility
# =========================================================================
s = slide()
chrome(s, "REPRODUCIBILITY", "재현성과 감사 — 이 연구의 축")
text(s, MARGIN, Inches(1.9), CONTENT_W, Inches(0.6),
     [[("MP4MR도 warehouse류 시스템도 갖지 않은 것: 같은 입력에 같은 결과를 보장하고, "
        "제3자가 판정을 재계산할 수 있게 만든다", dict(size=13, color=INK))]], leading=1.5)
items = [
    ("통합 event_log  ", "TurnAudit · ExecutionAudit · CheckpointAudit 를 발생 순서대로 — list 순서가 유일한 진실 원천"),
    ("4종 hash  ", "graph_hash · scene_hash · patch_hash · pre_state_hash — 모든 판정에 스탬프"),
    ("gold pre-commit  ", "모든 평가 정답을 LLM 호출 전에 커밋 — P12 held-out은 첫 결과를 본 뒤 별도로 다시 커밋"),
    ("결정론적 layout  ", "같은 graph → 같은 RenderSpec. 발표 그림도 UI와 같은 renderer가 산출"),
    ("contract-first  ", "계약을 코드보다 먼저 커밋 — 결정 이력 D-001~D-070"),
]
iy = Inches(2.7)
for lead, rest in items:
    rect(s, MARGIN, iy, Inches(0.14), Inches(0.14), DET, shape=MSO_SHAPE.OVAL)
    text(s, MARGIN + Inches(0.34), iy - Inches(0.06), CONTENT_W - Inches(0.4), Inches(0.7),
         [[(lead, dict(size=13, bold=True, color=INK)),
           (rest, dict(size=12.5, color=MUTED))]], leading=1.3)
    iy += Inches(0.66)

# =========================================================================
# 12 — scope / limits
# =========================================================================
s = slide()
chrome(s, "SCOPE", "주장하지 않는 것")
text(s, MARGIN, Inches(1.72), CONTENT_W, Inches(0.4),
     [[("발표에서 이 목록을 그대로 말한다 — 계약이 각 항목을 문서로 고정한다",
        dict(size=12, color=MUTED))]])
rect(s, MARGIN, Inches(2.25), CONTENT_W, Inches(4.55), RGBColor(0xF6, 0xE9, 0xE2),
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
vline(s, MARGIN, Inches(2.45), Inches(4.15), LIMIT, 3.0)
lims = [
    ("자동 화재 인식이 아니다.", " FIRE_DETECTED는 사전 고정 latent fixture가 정찰 완료 뒤 공개하는 결정론적 observation"),
    ("실제 telemetry가 아니다.", " 2D 이동은 discrete-event schedule의 kinematic 재생. 동역학·가속을 모사하지 않는다"),
    ("MP4MR·LSMP보다 우수하다고 말하지 않는다.", " 동일 조건 직접 비교를 하지 않았다 — CBBA vs BP도 마찬가지"),
    ("전역 최소 reset이 아니다.", " selective release는 결정론적 정책이며 최적 재계획을 증명하지 않는다"),
    ("Validator가 자연어 의미 충실도를 보장하지 않는다.", " Validator는 schema와 도메인 invariant만 판정 — 충실도는 precision/recall이 잰다"),
    ("소표본이다.", " 단일 scene · 단일 모델 스냅샷 · 단일 실행. 통계적 일반화를 하지 않는다"),
]
ly = Inches(2.5)
for a, b in lims:
    text(s, MARGIN + Inches(0.36), ly, CONTENT_W - Inches(0.7), Inches(0.66),
         [[(a, dict(size=12.5, bold=True, color=INK)),
           (b, dict(size=12, color=MUTED))]], leading=1.25)
    ly += Inches(0.66)

# =========================================================================
# 13 — conclusion
# =========================================================================
s = slide()
chrome(s, "CONCLUSION", "자연어가 임무 전 주기의 인터페이스다")
text(s, MARGIN, Inches(2.0), CONTENT_W, Inches(1.2),
     [[("계획할 때도, 실행 중 재계획할 때도 — 같은 5종 어휘, 같은 grounding · ",
        dict(size=15, color=INK)),
       ("같은 결정론적 검증", dict(size=15, color=INK, bold=True))],
      [("LLM은 말을 구조로 바꾸고, Validator와 CBBA가 그 구조를 검증하고 할당한다",
        dict(size=15, color=MUTED))]], leading=1.5)
text(s, MARGIN, Inches(3.7), Inches(3.2), Inches(0.4),
     [[("확인된 것", dict(mono=True, size=10, bold=True, color=DET, spacing=1.2))]])
text(s, MARGIN, Inches(4.05), Inches(6.0), Inches(2.0),
     bullets([
         "제약된 LLM 출력을 결정론적으로 검증 가능",
         "이종 UAV/UGV에 무위반 할당 · 실행",
         "실행 중 checkpoint에서 영향받은 task만 재할당",
         "명령이 달라지면 대응 그래프·경로가 달라진다 (held-out 8/8)",
     ], size=12.5), leading=1.35, space_after=5)
text(s, EMU_W / 2 + Inches(0.3), Inches(3.7), Inches(3.2), Inches(0.4),
     [[("후속", dict(mono=True, size=10, bold=True, color=LIMIT, spacing=1.2))]])
text(s, EMU_W / 2 + Inches(0.3), Inches(4.05), Inches(5.4), Inches(2.0),
     bullets([
         "platform 통합 (P7) — ROS2 / Gazebo 실행 backend",
         "recheck 계열 어휘 — 임의 시각 재계획",
         "다중 scene · 반복 실행으로 분산 측정",
         "더 약한 모델 · glossary 제거 조건",
     ], size=12.5), leading=1.35, space_after=5)

# ---- save --------------------------------------------------------------
out = Path(__file__).parent / "LLM-MRTA_구조.pptx"
prs.save(out)
print(f"wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
